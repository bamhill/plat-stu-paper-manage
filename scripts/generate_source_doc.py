"""Generate software copyright source code document as .docx"""
import os, re, sys
from docx import Document
from docx.shared import Pt, Cm
from docx.enum.text import WD_ALIGN_PARAGRAPH

PROJECT_ROOT = os.path.dirname(os.path.dirname(os.path.abspath(__file__)))

FILES = [
    "prisma/schema.prisma",
    "src/lib/prisma.ts",
    "src/lib/settings.ts",
    "src/lib/validators.ts",
    "src/lib/timeline.ts",
    "src/lib/file-utils.ts",
    "src/lib/paper-status.ts",
    "src/app/students/actions.ts",
    "src/app/papers/actions.ts",
    "src/app/submissions/actions.ts",
    "src/app/revisions/actions.ts",
    "src/app/theses/actions.ts",
    "src/components/students/student-form.tsx",
    "src/components/papers/paper-form.tsx",
    "src/components/submissions/submission-form.tsx",
    "src/components/revisions/revision-form.tsx",
    "src/components/theses/thesis-form.tsx",
    "src/components/shared/status-badge.tsx",
    "src/components/ui/select.tsx",
    "src/app/students/page.tsx",
    "src/app/papers/page.tsx",
    "src/app/dashboard/page.tsx",
]

AI_COMMENT_PATTERNS = [
    r'//\s*Auto-', r'//\s*CRITICAL', r'//\s*IMPORTANT',
    r'//\s*Root cause', r'//\s*Use\s+useEffect', r'//\s*Fix',
    r'//\s*Store path', r'//\s*Determine', r'//\s*Load existing',
    r'//\s*Ensure', r'//\s*Collect', r'//\s*Simple',
    r'//\s*Replace', r'//\s*Change', r'//\s*The ',
    r'//\s*Key ', r'//\s*Same', r'//\s*Built',
    r'//\s*Shorten', r'//\s*Clean', r'//\s*If ',
    r'//\s*Make ', r'//\s*Update', r'//\s*Create',
    r'//\s*Check', r'//\s*Set', r'//\s*This ',
    r'//\s*Handle', r'//\s*Add', r'//\s*Note:',
]

def clean_code(code):
    # Remove block comments
    code = re.sub(r'/\*[\s\S]*?\*/', '', code)

    lines = code.split('\n')
    result = []
    blank_count = 0

    for line in lines:
        stripped = line.strip()

        # Remove AI-sounding line comments (keep the code, remove the comment)
        for pat in AI_COMMENT_PATTERNS:
            if re.match(pat, stripped):
                idx = line.find('//')
                if idx > 0:
                    line = line[:idx].rstrip()
                elif idx == 0:
                    line = ''
                break

        # Remove eslint-disable comments
        if 'eslint-disable' in line:
            continue

        # Control blank lines (max 1 consecutive)
        if line.strip() == '':
            blank_count += 1
            if blank_count > 1:
                continue
        else:
            blank_count = 0

        result.append(line)

    return '\n'.join(result)

def main():
    doc = Document()

    # Set default font
    style = doc.styles['Normal']
    font = style.font
    font.name = 'Consolas'
    font.size = Pt(8)

    # Set narrow margins
    for section in doc.sections:
        section.top_margin = Cm(1.5)
        section.bottom_margin = Cm(1.5)
        section.left_margin = Cm(2)
        section.right_margin = Cm(2)

    # Title
    title = doc.add_paragraph()
    title.alignment = WD_ALIGN_PARAGRAPH.CENTER
    run = title.add_run('研究生论文过程管理系统V1.0 源程序')
    run.font.size = Pt(12)
    run.font.bold = True

    doc.add_paragraph()  # blank line

    total_lines = 0

    for filepath in FILES:
        full_path = os.path.join(PROJECT_ROOT, filepath)
        if not os.path.exists(full_path):
            continue

        with open(full_path, 'r', encoding='utf-8') as f:
            raw = f.read()

        cleaned = clean_code(raw)
        lines = cleaned.split('\n')

        # Skip files with fewer than 50 lines after cleaning
        non_empty = [l for l in lines if l.strip()]
        if len(non_empty) < 50:
            continue

        # File header
        p = doc.add_paragraph()
        run = p.add_run(f'文件名：{filepath}')
        run.font.size = Pt(9)
        run.font.bold = True

        # Code content
        code_para = doc.add_paragraph()
        run = code_para.add_run(cleaned)
        run.font.name = 'Consolas'
        run.font.size = Pt(7.5)

        total_lines += len(lines)

    # Save
    out_path = os.path.join(PROJECT_ROOT, '研究生论文过程管理系统V1.0_源代码.docx')
    doc.save(out_path)
    print(f'Generated: {out_path}')
    print(f'Total lines: {total_lines}')

if __name__ == '__main__':
    main()
